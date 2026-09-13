import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import csv
import datetime

columns = [
    ("Submission_ID", "Unique identifier for this intake submission"),
    ("Submitted_At", "Timestamp of form submission"),
    ("Account_Type", "Personal Portfolio or Agency / Studio"),
    ("Client_or_Brand_Name", "Name of the client, creator, or agency"),
    ("Professional_Role_or_Niche", "Personal role headline or agency market niche"),
    ("Discipline_or_Team_Size", "Personal discipline or agency headcount"),
    ("Availability_or_Engagement_Model", "Booking status or client engagement model"),
    ("Location_or_Studio_Model", "Personal location or agency studio model"),
    ("Resume_or_Project_Minimum", "Resume / CV link or agency minimum engagement budget"),
    ("Official_Email", "Primary client contact email"),
    ("Phone_WhatsApp", "Contact telephone or WhatsApp number"),
    ("Primary_Brand_Color", "Primary brand color in HEX"),
    ("Accent_Brand_Color", "Accent brand color in HEX"),
    ("Typography_Fonts", "Brand typography preferences"),
    ("Aesthetic_Tone", "Design aesthetic & visual tone"),
    ("Years_Experience", "Years of industry experience"),
    ("Clients_Served", "Total clients served metric"),
    ("Projects_Completed", "Total projects delivered metric"),
    ("Brand_Story_Mission", "Brand story / mission statement / personal bio"),
    ("Logo_Asset_Link", "Google Drive / Dropbox link to official logo vectors/PNGs"),
    ("Case_Study_1_Name", "Flagship Project 1 Title"),
    ("Case_Study_1_Problem", "Problem addressed in Project 1"),
    ("Case_Study_1_Solution", "Execution solution in Project 1"),
    ("Case_Study_1_Result", "Quantifiable impact / ROI in Project 1"),
    ("Case_Study_1_Drive_Link", "Google Drive link to Project 1 high-res images"),
    ("Case_Study_2_Name", "Flagship Project 2 Title"),
    ("Case_Study_2_Problem", "Problem addressed in Project 2"),
    ("Case_Study_2_Solution", "Execution solution in Project 2"),
    ("Case_Study_2_Result", "Quantifiable impact / ROI in Project 2"),
    ("Case_Study_2_Drive_Link", "Google Drive link to Project 2 high-res images"),
    ("Case_Study_3_Summary", "Project 3 overview (optional)"),
    ("Case_Study_4_Summary", "Project 4 overview (optional)"),
    ("Services_Offered", "List of core services and starting prices"),
    ("Client_Testimonials", "Verified client reviews and permission status"),
    ("Frequently_Asked_Questions", "Custom Q&As provided"),
    ("Physical_Address", "Business or remote address"),
    ("Preferred_Contact_Method", "WhatsApp, Email, Phone, or Calendly"),
    ("Business_Hours", "Operating hours and timezone"),
    ("Instagram", "Instagram profile link or handle"),
    ("LinkedIn", "LinkedIn URL"),
    ("Twitter_X", "X / Twitter URL or handle"),
    ("Behance_Dribbble", "Portfolio profile link"),
    ("Founder_Name", "Personal founder or creator name"),
    ("Founder_Title", "Professional title / designation"),
    ("Founder_Headshot_Link", "Drive link to high-res headshot"),
    ("Founder_Bio", "Executive or solo creator bio"),
    ("Case_Study_Architecture", "Structure choice (External / Dedicated / Hybrid)"),
    ("Pages_Required", "List of website pages to be designed"),
    ("Existing_Site_Status", "Brand New Build or Redesign"),
    ("Existing_Site_URL", "Current website URL if redesign"),
    ("Domain_Status", "Owned or assistance needed"),
    ("Target_Launch_Date", "Preferred deadline"),
    ("Reference_Websites", "Inspiration websites & design notes"),
    ("Copy_Status", "Ready or copywriting service required"),
    ("Revision_Policy_Agreed", "Terms acknowledged (Yes/No)"),
    ("Authorized_Signatory", "Full legal signature name"),
    ("Agreement_Date", "Date agreement executed")
]

wb = openpyxl.Workbook()

# Sheet 1: Master Submissions
ws = wb.active
ws.title = "Client Submissions"

header_fill = PatternFill(start_color="0A0D11", end_color="0A0D11", fill_type="solid")
header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
thin_border = Border(
    left=Side(style='thin', color='E2E8F0'),
    right=Side(style='thin', color='E2E8F0'),
    top=Side(style='thin', color='E2E8F0'),
    bottom=Side(style='thin', color='E2E8F0')
)

headers = [c[0] for c in columns]
ws.append(headers)

for col_num, col_name in enumerate(headers, 1):
    cell = ws.cell(row=1, column=col_num)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

# Add 2 sample demonstration rows (1 Personal, 1 Agency)
sample_personal = [
    "ALIE-INTAKE-2026-001",
    datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
    "Personal Portfolio",
    "Muhammad Ali",
    "Senior Product Designer & Framer Dev",
    "UI/UX & Product Design",
    "Available for Freelance & Sprints",
    "Lahore, Pakistan • Open to Global Remote",
    "https://drive.google.com/drive/folders/sample-cv",
    "ali@example.com",
    "+92 300 1234567",
    "#0A0D11",
    "#38BDF8",
    "Host Grotesk & Inter",
    "Futuristic & Cyber (Quest)",
    "5+ Years",
    "40+ Clients",
    "80+ Projects",
    "Crafting high-conversion digital flagships and interactive web systems for ambitious brands.",
    "https://drive.google.com/drive/folders/logo-sample",
    "Lumina Health Brand & App",
    "Outdated UI with sub-1% mobile conversions",
    "Complete Framer design sprint & sub-second performance upgrade",
    "+140% checkout conversion increase, $420k generated in Q1",
    "https://drive.google.com/drive/folders/lumina-images",
    "Fintech SaaS Mobile Dashboard",
    "Complex analytics interface with 60% user churn",
    "Redesigned intuitive progressive disclosure cards",
    "Cut activation time by 48%, Featured on ProductHunt",
    "https://drive.google.com/drive/folders/fintech-images",
    "Optional 3rd Project",
    "Optional 4th Project",
    "1. Custom Portfolio Website (PKR 65,000 / $350); 2. Brand Identity System (PKR 40,000 / $200)",
    "Hamza Tariq (CyberMatrix): 'ALIE CREATIVES transformed our digital image completely.' [Permission: Yes]",
    "Q1: Turnaround 5-7 days; Q2: 50% upfront deposit; Q3: Assets & copy needed; Q4: 3 revisions included; Q5: Tech & Creative niches",
    "Lahore, Pakistan (Remote Global)",
    "WhatsApp",
    "Mon-Sat, 10:00 AM - 7:00 PM PKT",
    "@ali_creatives",
    "https://linkedin.com/in/alie-creatives",
    "@alie_creatives",
    "https://behance.net/alie_creatives",
    "Muhammad Ali",
    "Creative Director & Founder",
    "https://drive.google.com/drive/folders/headshot-sample",
    "Bespoke designer specializing in high-performance digital experiences.",
    "Option B — Dedicated Case Study Pages",
    "Home, About, Services, Portfolio, Testimonials, FAQ, Contact",
    "Brand New Build",
    "",
    "Already Own Domain",
    "2026-10-01",
    "1. https://linear.app (Clean dark mode); 2. https://stripe.com (Smooth micro-interactions)",
    "Website copy is ready",
    "Yes — Agreed to Revision Policy",
    "Muhammad Ali",
    datetime.date.today().strftime("%Y-%m-%d")
]

sample_agency = [
    "ALIE-INTAKE-2026-002",
    datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
    "Agency / Studio",
    "Apex Studio LLC",
    "B2B SaaS & Tech Startups",
    "Growing Studio (6–15 team members)",
    "Fixed-Scope Fast Sprints (5–10 days)",
    "Physical HQ in Lahore + Remote Team",
    "150,000 – 400,000 PKR / $800 – $2,000",
    "contact@apexstudio.io",
    "+92 321 7654321",
    "#0F172A",
    "#A855F7",
    "Plus Jakarta Sans & Inter",
    "Bold & High-Contrast",
    "7+ Years",
    "120+ Clients",
    "250+ Delivered",
    "We architect full-lifecycle digital transformations for venture-backed SaaS startups.",
    "https://drive.google.com/drive/folders/apex-logo-vectors",
    "CloudGrid Infrastructure Cloud Portal",
    "Low enterprise trial signups due to dense documentation",
    "Built modern marketing flagship with interactive feature explorer",
    "+210% enterprise demo requests in 60 days",
    "https://drive.google.com/drive/folders/cloudgrid-shots",
    "Vertex Financial Banking Suite",
    "Slow legacy portal with high latency",
    "Next.js design system with sub-second page transitions",
    "Over $15M in deposits processed in first 90 days",
    "https://drive.google.com/drive/folders/vertex-shots",
    "",
    "",
    "1. Full Enterprise Flagship Website; 2. Product Design Sprint; 3. Visual Identity & Brand Book",
    "Sarah Jenkins (VP at SaaSFlow): 'World class design execution delivered ahead of schedule.' [Permission: Yes]",
    "Q1: 7-10 day enterprise sprints; Q2: 50/50 milestone billing; Q3: Brief & logo vectors; Q4: 3 revision rounds; Q5: B2B tech",
    "Gulberg III, Lahore, Pakistan",
    "WhatsApp",
    "Mon-Fri, 9:00 AM - 6:00 PM PKT",
    "@apexstudio",
    "https://linkedin.com/company/apex-studio",
    "@apex_digital",
    "https://behance.net/apexstudio",
    "Zayn Malik",
    "Managing Partner & Design Director",
    "https://drive.google.com/drive/folders/zayn-headshot",
    "10+ years directing digital experiences for international tech brands.",
    "Option C — Hybrid (Some linked, some built)",
    "Home, About, Services, Portfolio, Testimonials, FAQ, Contact",
    "Complete Redesign",
    "https://old-apexsite.com",
    "Already Own Domain",
    "2026-11-15",
    "1. https://vercel.com (Minimal contrast); 2. https://raycast.com (Dark obsidian theme)",
    "We need ALIE CREATIVES to write our website copy",
    "Yes — Agreed to Revision Policy",
    "Zayn Malik",
    datetime.date.today().strftime("%Y-%m-%d")
]

ws.append(sample_personal)
ws.append(sample_agency)

# Style data rows
data_font = Font(name="Segoe UI", size=10)
for row in ws.iter_rows(min_row=2, max_row=3, min_col=1, max_col=len(headers)):
    for cell in row:
        cell.font = data_font
        cell.border = thin_border
        cell.alignment = Alignment(vertical="center")

# Auto-adjust column widths
for col in ws.columns:
    max_len = max(len(str(cell.value or '')) for cell in col)
    col_letter = get_column_letter(col[0].column)
    ws.column_dimensions[col_letter].width = min(max(max_len + 3, 14), 45)

# Sheet 2: Schema & Instructions
ws2 = wb.create_sheet(title="Field Dictionary & Setup")
ws2.append(["Field Column Name", "Field Description", "Recommended Data Type / Format"])
ws2.cell(row=1, column=1).fill = header_fill
ws2.cell(row=1, column=1).font = header_font
ws2.cell(row=1, column=2).fill = header_fill
ws2.cell(row=1, column=2).font = header_font
ws2.cell(row=1, column=3).fill = header_fill
ws2.cell(row=1, column=3).font = header_font

for col_name, col_desc in columns:
    ws2.append([col_name, col_desc, "Text / String / URL"])

ws2.column_dimensions['A'].width = 36
ws2.column_dimensions['B'].width = 65
ws2.column_dimensions['C'].width = 30

# Save Excel file
excel_path = r"d:\ALI_Creatives\ALIE_CREATIVES_Client_Intake_Database.xlsx"
wb.save(excel_path)
print(f"Excel file saved successfully to {excel_path}")

# Save CSV version as well
csv_path = r"d:\ALI_Creatives\ALIE_CREATIVES_Client_Intake_Database.csv"
with open(csv_path, mode="w", newline="", encoding="utf-8-sig") as f:
    writer = csv.writer(f)
    writer.writerow(headers)
    writer.writerow(sample_personal)
    writer.writerow(sample_agency)
print(f"CSV file saved successfully to {csv_path}")
